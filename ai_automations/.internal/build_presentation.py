from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(11, 21, 39)
SURFACE = RGBColor(20, 37, 59)
SURFACE2 = RGBColor(28, 49, 72)
WHITE = RGBColor(245, 250, 255)
MUTED = RGBColor(155, 181, 204)
CYAN = RGBColor(89, 219, 213)
AMBER = RGBColor(249, 192, 109)
RED = RGBColor(255, 132, 136)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def shape(slide, x, y, w, h, color, radius=False):
    obj = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = color
    obj.line.fill.background()
    if radius:
        obj.adjustments[0] = 0.14
    return obj


def text(slide, content, x, y, w, h, size=22, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = Inches(0.02)
    box.text_frame.margin_right = Inches(0.02)
    box.text_frame.margin_top = Inches(0.02)
    box.text_frame.margin_bottom = Inches(0.02)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for idx, line in enumerate(content.split("\n")):
        p = box.text_frame.paragraphs[0] if idx == 0 else box.text_frame.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(6)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
    return box


def base(section, title, subtitle="", tag="DEMO  /  проверяемые правила"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape(slide, 0, 0, 13.333, 7.5, NAVY)
    shape(slide, 0.6, 0.62, 0.07, 0.38, CYAN)
    text(slide, section.upper(), 0.82, 0.62, 8.8, 0.3, 11, CYAN, True)
    text(slide, title, 0.8, 1.12, 11.7, 0.65, 33, WHITE, True)
    if subtitle:
        text(slide, subtitle, 0.82, 1.88, 11.7, 0.7, 16, MUTED)
    shape(slide, 0.8, 6.84, 11.75, 0.012, SURFACE2)
    text(slide, "REPKA   /   AI AUTOMATION", 0.82, 6.98, 4.5, 0.22, 9, MUTED, True)
    text(slide, tag, 6.7, 6.97, 5.3, 0.25, 9, MUTED, align=PP_ALIGN.RIGHT)
    text(slide, f"{len(prs.slides):02d}", 12.48, 6.97, 0.28, 0.23, 9, CYAN, True)
    return slide


def card(slide, x, y, w, h, heading, body, accent=CYAN):
    shape(slide, x, y, w, h, SURFACE, True)
    shape(slide, x + 0.23, y + 0.25, 0.07, 0.38, accent)
    text(slide, heading, x + 0.4, y + 0.22, w - 0.6, 0.44, 20, WHITE, True)
    text(slide, body, x + 0.4, y + 0.82, w - 0.8, h - 1.05, 16, MUTED)


def metric(slide, x, y, w, number, label, accent=CYAN):
    shape(slide, x, y, w, 1.56, SURFACE, True)
    text(slide, number, x + 0.27, y + 0.18, w - 0.5, 0.62, 34, accent, True)
    text(slide, label, x + 0.27, y + 0.9, w - 0.5, 0.5, 14, MUTED)


s = base("01 / Видение", "Repka AI Automation", "От данных Odoo к понятным сигналам и управляемым действиям")
shape(s, 0.8, 2.9, 7.3, 2.72, SURFACE, True)
text(s, "Меньше ручного поиска.\nБольше времени на решения.", 1.18, 3.38, 6.6, 1.6, 30, WHITE, True)
shape(s, 8.5, 2.9, 4.03, 2.72, SURFACE2, True)
text(s, "4 рабочих DEMO\n19 идей для пилота", 8.9, 3.4, 3.35, 1.43, 24, CYAN, True)
text(s, "Без подключений к production • без ML/LLM-обещаний", 0.85, 6.15, 11.5, 0.35, 15, MUTED)

s = base("02 / Проблема", "Ценные данные — разрозненные решения", "Сегодня для ответа на простой вопрос приходится открывать несколько экранов")
card(s, 0.82, 2.82, 3.73, 2.85, "Менеджер", "Собирает историю заказов клиента и ищет тех, кому стоит уделить внимание.")
card(s, 4.79, 2.82, 3.73, 2.85, "Логист", "Сверяет статусы документов и выясняет, какой заказ требует проверки.", AMBER)
card(s, 8.76, 2.82, 3.73, 2.85, "Руководитель", "Сопоставляет операционные сигналы без единой сводки.", RED)

s = base("03 / Возможность", "От поиска к фокусной проверке", "Автоматизируем подготовку контекста, не передаём машине право принимать решения")
card(s, 0.82, 2.82, 3.73, 2.9, "01  Собрать", "История клиента, суммы, статусы — из согласованных источников.")
card(s, 4.79, 2.82, 3.73, 2.9, "02  Объяснить", "Показать причину сигнала и исходные факты, не скрывая ограничений.", AMBER)
card(s, 8.76, 2.82, 3.73, 2.9, "03  Подтвердить", "Менеджер или логист проверяет факт и выбирает действие.")

s = base("04 / Данные Repka", "Основа уже есть в структуре Odoo", "Наличие моделей подтверждено кодом; полнота и качество реальных записей пока не подтверждены")
card(s, 0.82, 2.82, 5.69, 1.64, "Клиент + заказ", "res.partner  /  ufs.base.order")
card(s, 6.7, 2.82, 5.79, 1.64, "Экспедирование + затраты", "ufs.forwarder.order  /  подзаказы", AMBER)
card(s, 0.82, 4.64, 5.69, 1.64, "Документы + проверки", "ufs.transport.document  /  ufs.ai.document.check")
card(s, 6.7, 4.64, 5.79, 1.64, "Контекст платформы", "CRM  /  продажи  /  склад  /  закупки", AMBER)

s = base("05 / Архитектура", "Контролируемый контур автоматизации", "Сегодня: автономные правила на синтетике. Следующий шаг: согласованная read-only интеграция")
stages = [("Odoo / Repka", "Разрешённые поля"), ("AI layer", "Сейчас: правила"), ("Анализ", "Сигнал + причина"), ("Процесс", "Черновик / обзор"), ("Человек", "Решение")]
for i, (heading, desc) in enumerate(stages):
    x = 0.82 + i * 2.46
    shape(s, x, 3.02, 2.23, 2.1, SURFACE, True)
    text(s, heading, x + 0.19, 3.32, 1.86, 0.55, 18, CYAN, True)
    text(s, desc, x + 0.19, 4.13, 1.86, 0.62, 13, MUTED)
    if i < len(stages) - 1:
        text(s, "→", x + 2.25, 3.72, 0.22, 0.46, 22, AMBER, True)
text(s, "Без прямого доступа к БД • проверка ACL • человек подтверждает последствия", 0.85, 5.75, 11.7, 0.52, 17, WHITE)

s = base("06 / CRM", "Customer 360: кому уделить внимание", "Рабочий DEMO: понятное правило вместо недоказанного «скоринга вероятности»")
metric(s, 0.82, 2.76, 3.05, "149 дней", "после последнего заказа", AMBER)
card(s, 4.05, 2.76, 8.44, 2.45, "ООО Демо Север", "2 заказа • 1 неоплаченный • рекомендация: проверить необходимость контакта")
text(s, "Ввод → история перевозок     /     Расчёт → дни от последнего заказа     /     Выход → сигнал с причиной", 0.88, 5.53, 11.8, 0.76, 17, MUTED)

s = base("07 / Продажи", "Экономика заказа: увидеть отклонение", "Рабочий DEMO: сравниваются только суммы прямых затрат в одной валюте")
metric(s, 0.82, 2.76, 3.05, "−600 BYN", "результат DEMO-101", RED)
card(s, 4.05, 2.76, 8.44, 2.45, "Повод для финансовой проверки", "Выручка 5 000 BYN • прямые затраты 5 600 BYN • вывод: затраты выше выручки", RED)
text(s, "Не бухгалтерская прибыль: налоги, косвенные расходы и курсы не рассчитаны.", 0.88, 5.5, 11.8, 0.55, 17, MUTED)

s = base("08 / Документы", "Пакет без слепых зон", "Рабочий DEMO использует статусы проверок; не читает сканы и не заменяет OCR")
card(s, 0.82, 2.75, 5.7, 2.75, "Заказ DEMO-101", "CMR: blocked\nАкт: отсутствует", RED)
card(s, 6.75, 2.75, 5.74, 2.75, "Следующее действие", "Проверить замечания и оригиналы документов с логистом.", CYAN)
text(s, "Даже completed — только сигнал к проверке, а не юридическое заключение.", 0.88, 5.77, 11.8, 0.5, 17, MUTED)

s = base("09 / Руководителю", "Одна сводка — три точки внимания", "Рабочий DEMO: счётчики из синтетических срезов, не реальные показатели бизнеса")
metric(s, 0.82, 2.8, 3.7, "1", "клиент к проверке")
metric(s, 4.81, 2.8, 3.7, "1", "заказ: затраты > выручки", RED)
metric(s, 8.8, 2.8, 3.7, "1", "пакет документов", AMBER)
shape(s, 0.82, 4.73, 11.68, 1.36, SURFACE2, True)
text(s, "Результат: три конкретных вопроса ответственным, а не автоприказ системе", 1.13, 5.06, 11.1, 0.55, 21, WHITE, True)

s = base("10 / CRM и продажи", "Где потенциал после пилота", "Не реализовано: нужны история исходов, согласованные роли и данные для оценки качества", "FUTURE  /  не внедрено")
card(s, 0.82, 2.82, 3.73, 2.85, "Приоритет лидов", "Измерять на размеченных конверсиях, а не назначать «горячих» на глаз.")
card(s, 4.79, 2.82, 3.73, 2.85, "Прогноз спроса", "Строить интервалы только после проверки истории заказов.", AMBER)
card(s, 8.76, 2.82, 3.73, 2.85, "Подсказка контакта", "Создавать черновик; отправку подтверждает человек.")

s = base("11 / Склад и закупки", "Умное пополнение — при наличии фактов", "Стандартные модули есть, но подтверждённой истории остатков и сроков поставки нет", "FUTURE  /  не внедрено")
card(s, 0.82, 2.82, 5.68, 2.85, "Дефицит / избыточный запас", "Остатки + движения + сезонность → проверяемый прогноз риска.", AMBER)
card(s, 6.78, 2.82, 5.71, 2.85, "Качество поставщика", "Плановая и фактическая поставка → сигнал о задержке, не автозакупка.")

s = base("12 / Финансы и коммуникации", "Следующий контур автоматизации", "Идеи каталога; конфиденциальность и права важнее скорости внедрения", "FUTURE  /  не внедрено")
card(s, 0.82, 2.82, 5.68, 2.85, "Cash flow и расходы", "После сверки финансовых записей — сценарные прогнозы и поиск выбросов.", AMBER)
card(s, 6.78, 2.82, 5.71, 2.85, "Письма и задачи", "Только черновики и маршрутизация с разрешённым доступом и проверкой человека.")

s = base("13 / Путь к live", "Начать с данных, а не с обещаний", "Пилот начинается после решения открытых вопросов архитектуры", "ROADMAP  /  требует согласования")
for x, heading, body in [(0.82, "01  Проверить данные", "Объём, пропуски, валюты, роли"),
                         (4.79, "02  Замерить качество", "Размеченные случаи и ложные сигналы"),
                         (8.76, "03  Подключить безопасно", "Read-only API, ACL и аудит")]:
    card(s, x, 2.85, 3.73, 2.85, heading, body)

s = base("14 / Итог", "Repka + Odoo + AI", "Intelligent Business Automation — как целевая архитектура, не как текущий production-статус")
shape(s, 0.83, 2.72, 11.64, 2.9, SURFACE, True)
text(s, "Уже можно показать:\n4 воспроизводимых сценария на синтетических данных", 1.24, 3.05, 10.62, 1.3, 27, WHITE, True)
text(s, "Дальше: подтвердить данные → измерить пользу → согласовать защищённый live-пилот", 1.24, 4.75, 10.7, 0.52, 18, CYAN)

output = Path(__file__).resolve().parent.parent / "presentation" / "Repka_AI_Automation.pptx"
prs.save(output)
print(f"{output}: {len(prs.slides)} slides")
